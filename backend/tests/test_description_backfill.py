import asyncio
from io import BytesIO
from types import SimpleNamespace

from docx import Document
from sqlalchemy import create_engine, select

import app.db as db
from app.api import project_description
from app.db import Base, Chair, Department, Listing
from app.ingestion import description_backfill, llm_client
from app.ingestion.description_markdown import docx_to_markdown, html_to_markdown, inline_to_markdown
from app.ingestion.fetcher import Fetched
from app.ingestion.live import ingest_live
from app.ingestion.profile_sources import html_description_hash
from app.ingestion.registry import REGISTRY
from app.schemas import Status


def test_html_converter_keeps_offer_content_and_omits_page_chrome():
    html = b"""<body><nav>Other opportunity</nav><main><article>
        <h1>Robotics IDP</h1><p>Build a robot controller.</p>
        <ul><li>Python experience</li></ul><table><tr><th>Contact</th><td>Lab team</td></tr></table>
        </article></main><footer>Unrelated jobs</footer></body>"""
    markdown = html_to_markdown(html, title="Robotics IDP")
    assert "# Robotics IDP" in markdown
    assert "Build a robot controller." in markdown
    assert "- Python experience" in markdown
    assert "Contact | Lab team" in markdown
    assert "Other opportunity" not in markdown and "Unrelated jobs" not in markdown


def test_html_converter_selects_only_the_matching_offer_on_shared_page():
    html = b"""<main><div class='c-main'>
        <div class='frame'><h3>Robot IDP</h3><p>Build a robot controller.</p></div>
        <div class='frame'><h3>Climate IDP</h3><p>Model carbon emissions.</p></div>
        </div></main>"""
    markdown = html_to_markdown(html, title="Robot IDP", container_selector=".c-main > .frame")
    assert "Build a robot controller." in markdown
    assert "Climate IDP" not in markdown and "carbon emissions" not in markdown


def test_html_converter_rejects_unrelated_detail_page():
    html = b"<main><h1>Staff profile</h1><p>Laboratory contact details.</p></main>"
    try:
        html_to_markdown(html, title="Robotics IDP")
    except ValueError as error:
        assert "does not identify" in str(error)
    else:
        raise AssertionError("unrelated page must not become project Markdown")


def test_docx_converter_preserves_document_order():
    document = Document()
    document.add_heading("Robotics Project Study", level=1)
    document.add_paragraph("Build a robot controller.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Contact"
    table.cell(0, 1).text = "Lab team"
    output = BytesIO()
    document.save(output)
    markdown = docx_to_markdown(output.getvalue(), title="Robotics Project Study")
    assert markdown.index("# Robotics Project Study") < markdown.index("Build a robot controller.")
    assert markdown.index("Build a robot controller.") < markdown.index("Contact | Lab team")


def test_inline_converter_rejects_title_only_text():
    try:
        inline_to_markdown("Robotics IDP", title="Robotics IDP")
    except ValueError as error:
        assert "no substantive" in str(error)
    else:
        raise AssertionError("title-only text must not be treated as a description")


def test_png_converter_uses_image_input_and_returns_markdown(monkeypatch):
    seen = {}

    def parse(**kwargs):
        seen.update(kwargs)
        return SimpleNamespace(output_parsed=SimpleNamespace(markdown="# Robotics IDP\n\nBuild a robot."))

    monkeypatch.setattr(llm_client, "_client", lambda: SimpleNamespace(responses=SimpleNamespace(parse=parse)))
    markdown = llm_client.extract_image_markdown(b"\x89PNG\r\n\x1a\nexample", title="Robotics IDP")
    assert markdown.startswith("# Robotics IDP")
    assert seen["input"][0]["content"][1]["image_url"].startswith("data:image/png;base64,")
    assert seen["text_format"] is llm_client.ImageMarkdown


def test_backfill_persists_scoped_html_markdown_without_changing_listing_status(tmp_path, monkeypatch):
    test_engine = create_engine(f"sqlite:///{tmp_path / 'descriptions.db'}")
    monkeypatch.setattr(db, "_engine", test_engine)
    Base.metadata.create_all(test_engine)
    detail_url = "https://www.mdsi.tum.de/en/di-lab/projekte/ws26-robotics/"
    with db.session_factory()() as session:
        department = Department(slug="other-projects", name="Other Projects")
        session.add(department)
        session.flush()
        chair = Chair(slug="other-tum-data-innovation-lab", name="TUM Data Innovation Lab",
                      department_id=department.id)
        session.add(chair)
        session.flush()
        listing = Listing(chair_id=chair.id, stable_source_key="offer", slug="other-robotics",
                          reference_code="oth-001", title="Robotics project",
                          summary="Robotics project", normalized={"source_url": "https://www.mdsi.tum.de/en/di-lab/projects/",
                          "artifact_url": detail_url}, content_hash="source-hash", status=Status.active)
        session.add(listing)
        session.commit()

    content = b"<main><article><h1>Robotics project</h1><p>Build a robot controller for the lab.</p></article></main>"

    class FakeFetcher:
        async def fetch(self, url):
            assert url == detail_url
            return Fetched(url, 200, "text/html", content, "detail-hash", None, None)

    result = asyncio.run(description_backfill.run_description_backfill(fetcher=FakeFetcher()))
    assert result["converted"] == 1 and result["failed"] == 0
    with db.session_factory()() as session:
        row = session.scalar(select(Listing))
        assert row.status == Status.active
        assert row.normalized["description_source_url"] == detail_url
        assert row.normalized["description_content_hash"] == "detail-hash"
        markdown = row.normalized["description_markdown"]
        assert "Build a robot controller" in markdown
        assert row.normalized["description_hash"] == html_description_hash(detail_url, markdown)

    detail = project_description("other-robotics")
    assert detail.project.has_description
    assert "Build a robot controller" in detail.markdown
    assert asyncio.run(description_backfill.run_description_backfill(fetcher=FakeFetcher()))["selected"] == 0


def test_unapproved_child_url_is_never_fetched():
    target = description_backfill.DescriptionTarget(1, "other-tum-data-innovation-lab", "Offer", "",
        "https://www.mdsi.tum.de/en/di-lab/projects/", "https://example.test/unapproved", "hash")
    assert not description_backfill._allowed(target)


def test_failed_description_fetch_keeps_active_listing(tmp_path, monkeypatch):
    test_engine = create_engine(f"sqlite:///{tmp_path / 'failed-description.db'}")
    monkeypatch.setattr(db, "_engine", test_engine)
    Base.metadata.create_all(test_engine)
    with db.session_factory()() as session:
        department = Department(slug="other-projects", name="Other Projects")
        session.add(department)
        session.flush()
        chair = Chair(slug="other-tum-data-innovation-lab", name="TUM Data Innovation Lab",
                      department_id=department.id)
        session.add(chair)
        session.flush()
        session.add(Listing(chair_id=chair.id, stable_source_key="offer", slug="other-robotics",
                            reference_code="oth-001", title="Robotics project", summary="Robotics project",
                            normalized={"source_url": "https://www.mdsi.tum.de/en/di-lab/projects/",
                                        "artifact_url": "https://www.mdsi.tum.de/en/di-lab/projekte/ws26-robotics/"},
                            content_hash="source-hash", status=Status.active))
        session.commit()

    class FailingFetcher:
        async def fetch(self, _url):
            raise TimeoutError("detail page unavailable")

    result = asyncio.run(description_backfill.run_description_backfill(fetcher=FailingFetcher()))
    assert result["failed"] == 1 and result["converted"] == 0
    with db.session_factory()() as session:
        row = session.scalar(select(Listing))
        assert row.status == Status.active
        assert "description_markdown" not in row.normalized


def test_failed_detail_fetch_uses_substantive_listing_text(tmp_path, monkeypatch):
    test_engine = create_engine(f"sqlite:///{tmp_path / 'inline-fallback.db'}")
    monkeypatch.setattr(db, "_engine", test_engine)
    Base.metadata.create_all(test_engine)
    source_url = "https://www.mdsi.tum.de/en/di-lab/projects/"
    with db.session_factory()() as session:
        department = Department(slug="other-projects", name="Other Projects")
        session.add(department)
        session.flush()
        chair = Chair(slug="other-tum-data-innovation-lab", name="TUM Data Innovation Lab",
                      department_id=department.id)
        session.add(chair)
        session.flush()
        session.add(Listing(chair_id=chair.id, stable_source_key="offer", slug="other-robotics",
                            reference_code="oth-001", title="Robotics project",
                            summary="Develop a robot controller using camera input and validate it in the lab with the research team.",
                            normalized={"source_url": source_url,
                                        "artifact_url": "https://www.mdsi.tum.de/en/di-lab/projekte/ws26-robotics/"},
                            content_hash="source-hash", status=Status.active))
        session.commit()

    class FailingFetcher:
        async def fetch(self, _url):
            raise TimeoutError("detail page unavailable")

    result = asyncio.run(description_backfill.run_description_backfill(fetcher=FailingFetcher()))
    assert result["converted"] == 1 and result["fallback_inline"] == 1
    with db.session_factory()() as session:
        row = session.scalar(select(Listing))
        assert row.status == Status.active
        assert row.normalized["description_source_url"] == source_url
        assert "camera input" in row.normalized["description_markdown"]


def test_unchanged_crawl_keeps_backfilled_description_and_changed_offer_clears_it(tmp_path, monkeypatch):
    test_engine = create_engine(f"sqlite:///{tmp_path / 'recrawl.db'}")
    monkeypatch.setattr(db, "_engine", test_engine)
    Base.metadata.create_all(test_engine)
    adapter = REGISTRY[0]

    class FakeFetcher:
        def __init__(self, content):
            self.content = content

        async def fetch(self, url, **_kwargs):
            return Fetched(url, 200, "text/html", self.content, "top-hash", None, None)

    def page(description):
        return f"<article><h2>{adapter.name}: Project Study - Data</h2><p>{description}</p></article>".encode()

    asyncio.run(ingest_live((adapter,), FakeFetcher(page("First source description."))))
    with db.session_factory()() as session:
        listing = session.scalar(select(Listing))
        normalized = dict(listing.normalized)
        normalized["description_markdown"] = "# Project Study\n\nDetailed description."
        normalized["description_source_url"] = normalized["source_url"]
        normalized["description_hash"] = html_description_hash(normalized["source_url"],
                                                               normalized["description_markdown"])
        listing.normalized = normalized
        session.commit()

    asyncio.run(ingest_live((adapter,), FakeFetcher(page("First source description."))))
    with db.session_factory()() as session:
        assert session.scalar(select(Listing)).normalized["description_markdown"].endswith("Detailed description.")

    asyncio.run(ingest_live((adapter,), FakeFetcher(page("Changed source description."))))
    with db.session_factory()() as session:
        listing = session.scalar(select(Listing))
        assert listing.status == Status.active
        assert "description_markdown" not in listing.normalized
