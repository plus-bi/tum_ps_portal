from __future__ import annotations
from io import BytesIO
from pathlib import Path
from bs4 import BeautifulSoup


def html_text(data: bytes) -> str:
    soup = BeautifulSoup(data, "html.parser")
    for element in soup(["script", "style", "noscript"]): element.decompose()
    return "\n".join(s.strip() for s in soup.stripped_strings if s.strip())


def pdf_text(data: bytes, *, ocr=None) -> str:
    from .pdf_reader import USEFUL_TEXT_THRESHOLD, read_pdf, render_for_llm
    result = read_pdf(data)
    if result.native_text_length >= USEFUL_TEXT_THRESHOLD: return render_for_llm(result.pages).strip()
    if ocr is None: raise ValueError("PDF has no usable text and OCR is unavailable")
    return ocr(data)


def docx_text(data: bytes) -> str:
    from docx import Document
    doc = Document(BytesIO(data))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    cells = [c.text.strip() for table in doc.tables for row in table.rows for c in row.cells if c.text.strip()]
    return "\n".join(paragraphs + cells)


def extract(data: bytes, media_type: str, *, ocr=None) -> str:
    kind = media_type.split(";", 1)[0].casefold()
    if kind in {"text/html", "application/xhtml+xml"}: return html_text(data)
    if kind == "application/pdf": return pdf_text(data, ocr=ocr)
    if kind == "application/vnd.openxmlformats-officedocument.wordprocessingml.document": return docx_text(data)
    if kind.startswith("text/"): return data.decode("utf-8", errors="replace")
    raise ValueError(f"Unsupported source media type: {media_type}")
